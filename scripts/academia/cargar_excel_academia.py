#!/usr/bin/env python3
"""
Carga (o recarga) el Excel de la Academia en Supabase.

  python3 -I scripts/academia/cargar_excel_academia.py ACADEMIA.xlsx --seco
  python3 -I scripts/academia/cargar_excel_academia.py ACADEMIA.xlsx --llave ARCHIVO_CON_LA_LLAVE

--seco      solo lee el Excel e imprime conteos (no manda nada).
--llave     archivo con la llave temporal de carga (ver docs/migraciones/2026-10-13-academia-fase0-excel.md).

Todo se manda a la función public.academia_carga(), que es repetible: identifica cada registro
por hoja + teléfono (o nombre si no hay teléfono), nunca borra y no pisa lo que se edita en el Hub.
Este script no imprime nombres, teléfonos ni montos individuales.
"""
import argparse, collections, datetime, hashlib, json, re, sys, time, urllib.request, warnings
from pathlib import Path

warnings.filterwarnings('ignore')
import openpyxl
from openpyxl.utils import get_column_letter as letra

SUPABASE_URL = 'https://bomlszsnixazhnsegeik.supabase.co'
SUPABASE_KEY = 'sb_publishable_w85mi4Z_NunDmMQIcWGaaw_KqNSN2lq'   # clave publicable del Hub

# Hojas de alumnas → generación (se capturan tal cual; Academia corrige nombres y agrupaciones en el Hub)
GENERACIONES = [
    ('1ER GEN', 'gen1', '1ª generación'),
    ('2DO GEN', 'gen2', '2ª generación'),
    ('CHEF REPOSTERO EJECUTIVO - 3', 'gen3', '3ª generación'),
    ('CHEF REPOSTERO EJECUTIVO - 4', 'gen4', '4ª generación'),
    ('5TA GEN', 'gen5_mensual', '5ª generación mensual'),
    ('ACADEMIA 2025', 'ac2025', 'Academia 2025'),
    ('ACADEMIA 2026', 'ac2026', 'Academia 2026'),
    ('SOUS CHEF', 'sous_chef', 'Sous Chef'),
    ('HOBBY', 'hobby', 'Hobby'),
    ('REPOSTERO', 'repostero', 'Repostero'),
]
PROSPECTOS = [('PROSPECTOS', 'PROSPECTOS'), ('EXPO CINTERMEX', 'EXPO CINTERMEX'), ('INSCRITOS', 'INSCRITOS PENDIENTES DE PAGO')]
CONTENIDO = ['MODULOS', 'PLANES DE ACADEMIA', 'COPPYS DE ACADEMIA', 'CONVERSION', 'CLASES YOUTUBE', 'OBJECIONES MAS FRECUENTES']

MESES = 'ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|SEPTIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE|DIICIEMBRE'
COL_PAGO = re.compile(rf'^(ANTICIPO|ANT\s*\d*|INSC.*|MENS.*|{MESES})$')
COL_FECHA = re.compile(r'^(FECHA|FEHCA)$')
ENCUESTA = ('NEGOCIO', 'CUAL ES TU', 'COMO TE VES', 'EMPLEADOS', 'QUE QUIERES', 'LOGOTIPO', 'POR QUE ENTRAR',
            'PLATAFORMA', 'INGRESO A PLAT', 'COPPYS PARA', 'ACCECO', 'OBSERVACIONES')
BANCOS = ['HEY', 'COPPEL', 'BANORTE', 'BBVA', 'AZTECA', 'MERCADO', 'OXXO', 'EFECTIVO', 'SANTANDER', 'BANAMEX', 'SPIN']


# ---------- utilidades de limpieza ----------
def texto(v):
    """Valor de celda como texto, sin '.0' en enteros y fechas en ISO."""
    if v is None: return None
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat(sep=' ')
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    s = str(v).strip()
    return s or None

def telefono(v):
    """10 dígitos o None. Las fechas no son teléfonos y una lada no empieza en 0 o 1."""
    if v is None or isinstance(v, (datetime.datetime, datetime.date)): return None
    s = texto(v) or ''
    d = re.sub(r'\D', '', s.split('/')[0])
    d = d[-10:] if len(d) >= 10 else ''
    return d if d and d[0] in '23456789' else None

def normal(s):
    s = (s or '').upper()
    for a, b in zip('ÁÉÍÓÚÜÑ', 'AEIOUUN'): s = s.replace(a, b)
    return ' '.join(re.sub(r'[^A-Z0-9 ]', ' ', s).split())

def monto(t):
    if t is None: return None
    u = t.upper()
    if re.fullmatch(r'\$?\s*\d+(\.\d+)?', u.strip()): return float(re.sub(r'[^\d.]', '', u))
    m = re.search(r'\$\s?(\d{1,3}(?:,\d{3})+|\d+)', u) or re.search(r'(\d+)\s?\$', u)
    return float(m.group(1).replace(',', '')) if m else None

def pago(t, color, columna, nombre_col, fecha=None):
    m = monto(t); u = t.upper()
    if 'DEBE' in u: estado = 'debe'
    elif m is not None and m > 5000: estado = 'por_revisar'          # precio o monto con espacio comido
    elif 'PAGADO' in u or m is not None: estado = 'pagado'
    else: estado = 'por_revisar'
    return {'columna': columna, 'columna_excel': nombre_col, 'texto': t, 'fecha_texto': fecha,
            'monto': m, 'banco': next((b for b in BANCOS if b in u), None), 'color': color, 'estado': estado}

def color_de(c):
    rgb = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == 'solid' else None
    return rgb if isinstance(rgb, str) and rgb not in ('FFFFFFFF', '00000000') else None

def correo_valido(s):
    s = (s or '').strip().lower()
    return s if re.fullmatch(r'[^@\s]+@[^@\s]+\.[a-z]{2,}', s) else None

def nacimiento(v):
    if isinstance(v, datetime.datetime) and 1930 <= v.year <= 2012: return v.date().isoformat()
    return None

def decision(grad):
    u = normal(grad)
    if not u: return 'por_contactar'
    if u.startswith('GRADUADA'): return 'graduada'
    if u.startswith('KIT'): return 'kit'
    if u.startswith('GRAD 20') or u == 'SI': return 'se_gradua'
    if 'NO QUIERE' in u or 'NO ASISTIRA' in u: return 'no_quiere'
    if 'ESPERA' in u or 'SIGUIENTE' in u or 'PROXIMA' in u: return 'espera'
    if u.startswith('MSJ'): return 'contactada'
    return 'por_contactar'

def hoja(wb, prefijo):
    return next((w for w in wb.worksheets if w.title.startswith(prefijo)), None)

def fila_encabezado(ws, nombre='NOMBRE', hasta=5):
    return next((r for r in range(1, hasta + 1)
                 if any((texto(c.value) or '').upper() == nombre for c in ws[r])), None)

def clave_unica(vistas, base, fila):
    k = base if base not in vistas else f'{base}:f{fila}'
    vistas.add(k); return k


# ---------- capa 1 ----------
def celdas(wb_v, wb_f):
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value in (None, ''): continue
                f = wf.cell(c.row, c.column).value
                yield {'hoja': ws.title, 'fila': c.row, 'columna': c.column, 'ref': f'{letra(c.column)}{c.row}',
                       'valor': texto(c.value), 'tipo': c.data_type,
                       'formula': f if isinstance(f, str) and f.startswith('=') else None,
                       'color': color_de(c), 'comentario': c.comment.text if c.comment else None,
                       'link': c.hyperlink.target if c.hyperlink else None}


# ---------- alumnas ----------
def alumnas(wb):
    salida, resumen = [], {}
    for prefijo, gen, _ in GENERACIONES:
        ws = hoja(wb, prefijo)
        if ws is None: continue
        hr = fila_encabezado(ws)
        hdr = [(texto(c.value) or '').upper() for c in ws[hr]]
        hdr_crudo = [c.value for c in ws[hr]]
        ni = hdr.index('NOMBRE')
        fin_nombre = next((i for i in range(ni + 1, len(hdr)) if hdr[i]), ni + 4)
        cols_nombre = list(range(ni, min(fin_nombre, ni + 4)))
        def col(*nombres):
            return next((i for i, h in enumerate(hdr) if h in nombres), None)
        c_tel, c_tel2 = col('CELULAR', 'WPP'), col('TEL. 2', 'TEL 2')
        c_mat = 0 if hdr[0] == 'MATRICULA' else None
        c_plan = 0 if hdr[0] and hdr[0] != 'MATRICULA' and 0 not in cols_nombre else None
        pagos_cols = []
        for i, h in enumerate(hdr):
            es_mes = isinstance(hdr_crudo[i], datetime.datetime)
            if es_mes or COL_PAGO.match(h):
                fecha_col = i + 1 if i + 1 < len(hdr) and COL_FECHA.match(hdr[i + 1]) else None
                pagos_cols.append((i, texto(hdr_crudo[i]) or letra(i + 1), fecha_col))
        seccion, vistas, cont = None, set(), collections.Counter()
        for r in range(hr + 1, ws.max_row + 1):
            row = list(ws[r])
            vals = [c.value for c in row]
            if not any(v not in (None, '') for v in vals): continue
            partes = [texto(vals[i]) for i in cols_nombre if i < len(vals) and texto(vals[i]) and not telefono(vals[i])]
            nombre = ' '.join(partes).upper()
            if nombre == 'NOMBRE' or (partes and partes[0].upper() == 'NOMBRE'):
                cont['encabezado_repetido'] += 1; continue
            t = (telefono(vals[c_tel]) if c_tel is not None else None) or \
                (telefono(vals[c_tel2]) if c_tel2 is not None else None) or \
                next((telefono(vals[i]) for i in cols_nombre if i < len(vals) and telefono(vals[i])), None)
            llenas = sum(1 for i, v in enumerate(vals) if v not in (None, '') and i not in cols_nombre)
            # Separador = un solo texto en la columna del nombre y nada más ("DESERTORES", "TORREON AC 2026").
            # Un nombre con apellidos o con algún otro dato es alumna, aunque no tenga teléfono.
            if not t and len(partes) <= 1 and llenas == 0:
                if nombre: seccion = nombre[:120]; cont['separadores'] += 1
                continue
            if not nombre and not t:
                cont['sin_nombre_ni_telefono'] += 1; continue          # quedan solo en la capa 1
            def campo(*nombres):
                i = col(*nombres)
                return texto(vals[i]) if i is not None and i < len(vals) else None
            fila_excel, encuesta, redes, importes, comentarios = {}, {}, {}, {}, []
            for i, v in enumerate(vals):
                tv = texto(v)
                if tv is None: continue
                h = hdr[i] if i < len(hdr) else ''
                fila_excel[f'{letra(i + 1)}:{h}'] = tv
                if h.startswith(ENCUESTA): encuesta[h] = tv
                elif h in ('FACEBOOK', 'FB', 'INSTAGRAM'): redes[h] = tv
                elif h in ('RESTANTE', 'TOTAL', 'SUBTOTAL', 'LIQUIDO', 'MENSAULIDAD'): importes[h] = tv
                elif h == 'COMENTARIO': comentarios.append(tv)
            pagos = []
            for i, nom_col, fcol in pagos_cols:
                if i >= len(vals) or texto(vals[i]) is None: continue
                fecha = texto(vals[fcol]) if fcol is not None and fcol < len(vals) else None
                pagos.append(pago(texto(vals[i]), color_de(row[i]), i + 1, nom_col, fecha))
            grad = campo('GRADUACION', 'GRADUACIÓN')
            correo = campo('CORREO')
            nac_i = col('NACIMIENTO', 'F NACIMIENTO', 'F NACIMINETO')
            nac_v = vals[nac_i] if nac_i is not None and nac_i < len(vals) else None
            base = t or 'sintel:' + normal(nombre)
            salida.append({
                'gen': gen, 'clave': clave_unica(vistas, base, r), 'hoja': ws.title, 'fila': r, 'seccion': seccion,
                'matricula': texto(vals[c_mat]) if c_mat is not None else None,
                'nombre': nombre or '(SIN NOMBRE)', 'telefono': t,
                'telefono_2': telefono(vals[c_tel2]) if c_tel2 is not None else None,
                'ciudad': campo('CIUDAD'), 'vendedor': campo('VENDEDOR', 'PROMOTOR'),
                'plan': texto(vals[c_plan]) if c_plan is not None else None,
                'importes': importes, 'graduacion': grad, 'comentario': ' | '.join(comentarios) or None,
                'decision': decision(grad), 'correo': correo_valido(correo), 'nacimiento': nacimiento(nac_v),
                'privados': {'curp': campo('CURP'), 'direccion': campo('DIRECCION'), 'estado_civil': campo('EST CIVIL'),
                             'nacimiento_texto': texto(nac_v), 'correo_excel': correo,
                             'redes': redes, 'encuesta': encuesta, 'fila_excel': fila_excel},
                'pagos': pagos})
            cont['alumnas'] += 1; cont['sin_telefono'] += 0 if t else 1; cont['pagos'] += len(pagos)
        resumen[ws.title] = dict(cont)
    return salida, resumen


# ---------- graduación ----------
def tramos(ws, r_hdr):
    """Cada encabezado cubre desde su columna hasta antes del siguiente encabezado."""
    cols = [(c.column, (texto(c.value) or '').upper()) for c in ws[r_hdr] if texto(c.value)]
    return [(h, a, (cols[i + 1][0] if i + 1 < len(cols) else a + 1)) for i, (a, h) in enumerate(cols)]

def leer_tramo(ws, r, a, b):
    partes = [texto(ws.cell(r, c).value) for c in range(a, b)]
    return ' · '.join(p for p in partes if p) or None

def graduacion(wb):
    ws = hoja(wb, 'GRADUACION AC')
    def conceptos(c_txt, c_num, r_fin=50):
        lista, alumna, acomp = [], None, None
        for r in range(1, r_fin):
            n, v = texto(ws.cell(r, c_txt).value), ws.cell(r, c_num).value
            if n and isinstance(v, (int, float)):
                if n.upper().startswith('ALUMNO'): alumna = v
                elif n.upper().startswith('ACOMPA'): acomp = v
                else: lista.append({'concepto': n, 'precio': v})
        return lista, alumna, acomp
    def itinerario(c_ini, c_fin, r_fin=50):
        return [x for x in (leer_tramo(ws, r, c_ini, c_fin + 1) for r in range(1, r_fin)) if x]
    con_a, al_a, ac_a = conceptos(19, 20)     # S / T
    con_b, al_b, ac_b = conceptos(39, 40)     # AM / AN
    grads = [
        {'clave': '2026-08', 'nombre': 'Graduación agosto 2026', 'fechas_texto': texto(ws['A1'].value),
         'fecha_inicio': '2026-08-04', 'generaciones_texto': ' · '.join(filter(None, [texto(ws['I1'].value), texto(ws['K1'].value), texto(ws['M1'].value)])),
         'precio_alumna': al_a, 'precio_acompanante': ac_a, 'conceptos': con_a, 'itinerario': itinerario(1, 18), 'estado': 'pasada'},
        {'clave': '2027-08', 'nombre': 'Graduación agosto 2027', 'fechas_texto': texto(ws['U1'].value),
         'fecha_inicio': '2027-08-03', 'generaciones_texto': 'GEN 4, 5 Y 6 (según el título del Excel)',
         'precio_alumna': al_b, 'precio_acompanante': ac_b, 'conceptos': con_b, 'itinerario': itinerario(21, 38), 'estado': 'proxima'},
    ]
    # Bloques de listas (todas son de la graduación de agosto 2026)
    hdrs = [r for r in range(1, ws.max_row + 1) if (texto(ws.cell(r, 1).value) or '').upper().startswith('NOMBRE')]
    asistentes, kits, hotel = [], [], []
    vistas = {'a': set(), 'k': set(), 'h': set()}
    for i, r_hdr in enumerate(hdrs):
        fin = hdrs[i + 1] if i + 1 < len(hdrs) else ws.max_row + 1
        tr = tramos(ws, r_hdr)
        cab = {h: (a, b) for h, a, b in tr}
        primera = (texto(ws.cell(r_hdr, 1).value) or '').upper()
        etiqueta = texto(ws.cell(r_hdr - 1, 1).value)
        tipo = 'asistentes' if primera.startswith('NOMBRE ALUMN') else ('hotel' if 'SENCILLA' in cab else 'kits')
        for r in range(r_hdr + 1, fin):
            if not any(texto(ws.cell(r, c).value) for c in range(1, 41)): continue
            def t(h):
                return leer_tramo(ws, r, *cab[h]) if h in cab else None
            nombre = ' '.join(filter(None, (texto(ws.cell(r, c).value) for c in range(1, 5) if not telefono(ws.cell(r, c).value)))).upper()
            tel = telefono(ws.cell(r, 5).value)
            if not nombre and not tel: continue
            datos = {f'{letra(c)}:{(texto(ws.cell(r_hdr, c).value) or "").upper()}': texto(ws.cell(r, c).value)
                     for c in range(1, 41) if texto(ws.cell(r, c).value)}
            base = tel or 'sintel:' + normal(nombre)
            if tipo == 'asistentes':
                bloque = etiqueta if etiqueta and not telefono(ws.cell(r_hdr - 1, 5).value) else f'sin etiqueta (fila {r_hdr})'
                asistentes.append({'clave': clave_unica(vistas['a'], base, r), 'bloque': bloque, 'fila': r,
                    'nombre': nombre or '(SIN NOMBRE)', 'telefono': tel,
                    'invitados_texto': ' · '.join(filter(None, [t('INV'), t('CLUB'), t('INVITADO'), t('NOMBRE'), t('PARENTESCO')])) or None,
                    'filipina_texto': t('FILIPINA'), 'hotel_texto': t('HOTEL'), 'liquido_texto': t('LIQUIDO'),
                    'abonos': [x for x in (t(f'ANT {n}') for n in range(1, 6)) if x], 'datos': datos})
            elif tipo == 'kits':
                cf = cab.get('FILIPINA', (13, 14))[0]
                kits.append({'clave': clave_unica(vistas['k'], base, r), 'fila': r, 'nombre': nombre or '(SIN NOMBRE)',
                    'telefono': tel, 'generacion_texto': t('GEN'), 'direccion': t('DIRECCIÓN'),
                    'anticipo_texto': t('ANTICIPO'), 'filipina_texto': texto(ws.cell(r, cf).value),
                    'pago_texto': texto(ws.cell(r, cf + 1).value), 'envio_texto': t('ENVIO'),
                    'situacion_filipina': t('SITUACIÓN FILIPINA'), 'datos': datos})
            else:
                hotel.append({'clave': clave_unica(vistas['h'], base, r), 'fila': r, 'nombre': nombre or '(SIN NOMBRE)',
                    'telefono': tel, 'generacion_texto': t('GEN'), 'sencilla': t('SENCILLA'), 'doble': t('DOBLE'),
                    'triple': t('TRIPLE'), 'cuadruple': t('CUADRUPLE'), 'entrada': t('ENTRADA'), 'salida': t('SALIDA'),
                    'anticipo_texto': t('ANT IMP'), 'impuestos_texto': next((leer_tramo(ws, r, a, b) for h, a, b in tr if h.startswith('IMP')), None),
                    'total_texto': t('TOTAL'), 'datos': datos})
    return grads, asistentes, kits, hotel


# ---------- prospectos y contenido ----------
def prospectos(wb):
    salida = []
    for prefijo, origen in PROSPECTOS:
        ws = hoja(wb, prefijo)
        hr = fila_encabezado(ws)
        hdr = [(texto(c.value) or '').upper() for c in ws[hr]]
        ni = hdr.index('NOMBRE')
        cols_nombre = [i for i in range(ni, min(len(hdr), ni + 4)) if i == ni or not hdr[i] or hdr[i].startswith('APELLIDO')]
        c_tel = next((i for i, h in enumerate(hdr) if h in ('TELEFONO', 'WHATSAPP', 'CELULAR')), None)
        vistas = set()
        for r in range(hr + 1, ws.max_row + 1):
            vals = [c.value for c in ws[r]]
            if not any(v not in (None, '') for v in vals): continue
            nombre = ' '.join(texto(vals[i]) for i in cols_nombre if i < len(vals) and texto(vals[i]) and not telefono(vals[i])).upper()
            tel = (telefono(vals[c_tel]) if c_tel is not None and c_tel < len(vals) else None) or \
                  next((telefono(vals[i]) for i in cols_nombre if i < len(vals) and telefono(vals[i])), None)
            if not nombre and not tel: continue
            datos = {f'{letra(i + 1)}:{hdr[i] if i < len(hdr) else ""}': texto(v) for i, v in enumerate(vals) if texto(v)}
            salida.append({'origen': origen, 'clave': clave_unica(vistas, tel or 'sintel:' + normal(nombre), r), 'fila': r,
                           'nombre': nombre or '(SIN NOMBRE)', 'telefono': tel,
                           'ciudad': texto(vals[hdr.index('CIUDAD')]) if 'CIUDAD' in hdr and hdr.index('CIUDAD') < len(vals) else None,
                           'atendido': texto(vals[hdr.index('ATENDIDO')]) if 'ATENDIDO' in hdr and hdr.index('ATENDIDO') < len(vals) else None,
                           'datos': datos})
    return salida

def contenido(wb):
    for nombre in CONTENIDO:
        ws = hoja(wb, nombre)
        for row in ws.iter_rows():
            for c in row:
                if c.value in (None, ''): continue
                yield {'hoja': ws.title, 'fila': c.row, 'columna': c.column, 'texto': texto(c.value),
                       'link': c.hyperlink.target if c.hyperlink else None}


# ---------- envío ----------
def llamar(llave, accion, datos, intentos=3):
    cuerpo = json.dumps({'p_token': llave, 'p_accion': accion, 'p_datos': datos}, ensure_ascii=False).encode()
    req = urllib.request.Request(f'{SUPABASE_URL}/rest/v1/rpc/academia_carga', data=cuerpo, method='POST',
        headers={'apikey': SUPABASE_KEY, 'Authorization': f'Bearer {SUPABASE_KEY}', 'Content-Type': 'application/json'})
    for n in range(intentos):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return json.loads(r.read())
        except urllib.error.HTTPError as e:
            msg = e.read().decode()[:400]
            if n == intentos - 1 or e.code < 500: raise SystemExit(f'Error en {accion}: {e.code} {msg}')
        except urllib.error.URLError:
            if n == intentos - 1: raise
        time.sleep(2 ** (n + 1))

def lotes(lista, n):
    for i in range(0, len(lista), n): yield lista[i:i + n]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('excel'); ap.add_argument('--llave'); ap.add_argument('--seco', action='store_true')
    ap.add_argument('--archivo', help='id de una carga anterior del MISMO archivo: no vuelve a copiar las celdas')
    ap.add_argument('--desde', default='generaciones',
                    choices=['generaciones', 'alumnas', 'graduacion', 'prospectos', 'contenido'],
                    help='retomar una carga que se cortó a partir de este paso')
    a = ap.parse_args()
    ruta = Path(a.excel)
    md5 = hashlib.md5(ruta.read_bytes()).hexdigest()
    wb_v = openpyxl.load_workbook(ruta, data_only=True)
    wb_f = openpyxl.load_workbook(ruta, data_only=False)
    lista_celdas = list(celdas(wb_v, wb_f))
    lista_al, resumen_al = alumnas(wb_v)
    grads, asis, kits, hot = graduacion(wb_v)
    pros = prospectos(wb_v)
    cont = list(contenido(wb_v))
    resumen = {'md5': md5, 'celdas': len(lista_celdas), 'alumnas_por_hoja': resumen_al,
               'alumnas': len(lista_al), 'pagos': sum(len(x['pagos']) for x in lista_al),
               'asistentes_2026': len(asis), 'kits': len(kits), 'hotel': len(hot),
               'prospectos': dict(collections.Counter(p['origen'] for p in pros)), 'contenido': len(cont)}
    print(json.dumps(resumen, ensure_ascii=False, indent=1))
    if a.seco: return
    llave = Path(a.llave).read_text().strip()

    if a.archivo:
        arch = a.archivo
    else:
        arch = llamar(llave, 'archivo', {'nombre': ruta.name, 'md5': md5, 'notas': 'carga con cargar_excel_academia.py'})['archivo_id']
        n = 0
        for lote in lotes(lista_celdas, 2500):
            n += llamar(llave, 'celdas', {'archivo_id': arch, 'celdas': lote})['celdas']
        print('capa 1 · celdas:', n, '· archivo_id:', arch)
    pasos = ['generaciones', 'alumnas', 'graduacion', 'prospectos', 'contenido']
    hacer = set(pasos[pasos.index(a.desde):])
    if 'generaciones' in hacer:
        print('generaciones:', llamar(llave, 'generaciones', {'filas': [
            {'clave': g, 'nombre': nom, 'hoja': (hoja(wb_v, p).title if hoja(wb_v, p) else p), 'orden': i + 1}
            for i, (p, g, nom) in enumerate(GENERACIONES)]}))
    if 'alumnas' in hacer:
        tot = collections.Counter()
        for lote in lotes(lista_al, 150):
            tot.update(llamar(llave, 'alumnas', {'filas': lote}))
        print('alumnas:', dict(tot))
    if 'graduacion' in hacer:
        for g in grads:
            print('graduación', g['clave'], llamar(llave, 'graduacion', g))
        print('asistentes:', llamar(llave, 'asistentes', {'graduacion': '2026-08', 'filas': asis}))
        print('kits:', llamar(llave, 'kits', {'graduacion': '2026-08', 'filas': kits}))
        print('hotel:', llamar(llave, 'hotel', {'graduacion': '2026-08', 'filas': hot}))
    if 'prospectos' in hacer:
        tp = collections.Counter()
        for lote in lotes(pros, 300):
            tp.update(llamar(llave, 'prospectos', {'filas': lote}))
        print('prospectos:', dict(tp))
    tc = collections.Counter()
    for lote in lotes(cont, 1000):
        tc.update(llamar(llave, 'contenido', {'filas': lote}))
    print('contenido:', dict(tc))
    llamar(llave, 'cerrar_archivo', {'archivo_id': arch, 'resumen': resumen})
    print('archivo_id:', arch)

if __name__ == '__main__':
    main()
